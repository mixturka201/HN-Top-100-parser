import os
import requests
import time
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font

BASE_URL = "https://hacker-news.firebaseio.com/v0"
OUTPUT_FILE = "hn_top100.xlsx"
LIMIT = 100

TG_BOT_TOKEN = os.getenv("TG_BOT_TOKEN")
TG_CHAT_ID = os.getenv("TG_CHAT_ID")

def fetch_top_ids(limit=LIMIT):
    url = f"{BASE_URL}/topstories.json"
    response = requests.get(url, timeout=10)
    response.raise_for_status()
    return response.json()[:limit]

def fetch_item(item_id):
    url = f"{BASE_URL}/item/{item_id}.json"
    response = requests.get(url, timeout=10)
    response.raise_for_status()
    return response.json()

def collect_stories(ids):
    stories = []

    for i, item_id in enumerate(ids, 1):
        print(f"Loading {i}/{len(ids)}...")

        try:
            item = fetch_item(item_id)
        except requests.RequestException as e:
            print(f"Error: {e}")
            continue

        if item is None:
            continue
        
        stories.append({
            "title": item.get("title", ""),
            "author": item.get("by", ""),
            "score": item.get("score", 0),
            "comments": item.get("descendants", 0),
            "date": datetime.fromtimestamp(item.get("time", 0)).strftime("%Y-%m-%d %H:%M"),
            "url": item.get("url", f"https://news.ycombinator.com/item?id={item_id}"),
            })

        time.sleep(0.1)
    
    return stories

def save_to_excel(stories, filename=OUTPUT_FILE):
    wb = Workbook()
    ws = wb.active
    ws.title = "HN Top 100"

    headers = ["Title", "Author", "Score", "Comments", "Date", "URL"]
    ws.append(headers)

    for col in range(1, len(headers) + 1):
        ws.cell(row=1, column=col).font = Font(bold=True)

    for s in stories:
        ws.append([s["title"], s["author"], s["score"], s["comments"], s["date"], s["url"]])

    ws.column_dimensions["A"].width = 60
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 12
    ws.column_dimensions["E"].width = 18
    ws.column_dimensions["F"].width = 60

    for row in range(2, ws.max_row + 1):
        cell = ws.cell(row=row, column=6)
        cell.hyperlink = cell.value
        cell.font = Font(color="0000FF", underline="single")
    
    wb.save(filename)

def send_to_telegram(stories):
    if not TG_BOT_TOKEN or not TG_CHAT_ID:
        print("Telegram credentials not set. Skipping.")
        return

    lines = ["<b>HN Top 5:</b>\n"]

    for i, s in enumerate(stories[:5], 1):
        lines.append(f'{i}. <a href="{s["url"]}">{s["title"]}</a> ({s["score"]} points)')

    text = "\n".join(lines)
    url = f"https://api.telegram.org/bot{TG_BOT_TOKEN}/sendMessage"

    payload = {
        "chat_id": TG_CHAT_ID,
        "text": text,
        "parse_mode": "HTML",
        "disable_web_page_preview": True,
    }

    try:
        response = requests.post(url, json=payload, timeout=10)
        response.raise_for_status()
        print("Telegram message sent.")
    except requests.RequestException as e:
        print(f"Telegram error: {e}")

def main():
    print("Fetching top stories IDs...")
    ids = fetch_top_ids()

    print(f"Collecting {len(ids)} stories...")
    stories = collect_stories(ids)

    print("Sorting by score...")
    stories.sort(key=lambda x: x["score"], reverse=True)

    print(f"Saving to {OUTPUT_FILE}...")
    save_to_excel(stories)

    print(f"Done! {len(stories)} stories saved.")

if __name__ == "__main__":
    main()